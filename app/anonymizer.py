"""Detect and censor personal data in text, PDF, DOCX and TXT files.

Each sensitive word is replaced by a single '*'.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import pymupdf
import unicodedata
from docx import Document

import llm

LABEL_KEYWORDS = (
    "nome", "name", "nascimento", "birth", "sexo", "gender", "nif", "cartão", "cartao",
    "cidadão", "cidadao", "morada", "address", "telefone", "telemóvel", "phone", "tel",
    "email", "e-mail", "segurança social", "seguranca social", "social security",
    "estado civil", "marital", "filhos", "children", "religi", "ocupa", "profiss",
    "occupation", "hábitos", "habitos", "substânc", "substanc", "seguro", "insurance",
    "apólice", "apolice", "crédito", "credito", "credit", "rendimento", "salário",
    "salario", "salary", "income", "biom", "impress", "reconhecimento", "altura",
    "height", "peso", "weight", "sangu", "blood", "pressão arterial", "blood pressure",
    "crm", "passport", "passaporte", "iban", "referência", "referencia", "reference",
)
NAME_LABELS = ("nome", "name")
CHILD_LABELS = ("filhos", "children")

CAP = r"[A-ZÀ-Ý][A-Za-zÀ-ÿ'\-]+"
TITLE_RE = re.compile(rf"\b(?:Dr|Dra|Sr|Sra|Prof|Eng)\.?\s+({CAP}(?:\s+{CAP})*)")
PATIENT_RE = re.compile(rf"\b(?:paciente|patient)\s+({CAP}(?:\s+{CAP})*)")
STOP = {"Rua", "Avenida", "Lisboa", "Porto", "Centro", "Médico", "Medico"}

PATTERNS = [
    re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"),
    re.compile(r"\b[A-Z]{2}\d{2}(?:\s?\d{4}){3,7}\b"),
    re.compile(r"\b(?:\d{4}[ -]){3}\d{4}\b"),
    re.compile(r"(?i)\bCVV\s*\d{3,4}\b"),
    re.compile(r"(?<!\w)(?:\+|00)?\d{2,3}[ .-]?\d{3}[ .-]?\d{3}[ .-]?\d{3}\b"),
    re.compile(r"\b\d{2}/\d{2}/\d{4}\b|\b\d{4}-\d{2}-\d{2}\b"),
    re.compile(r"\b\d{8}-\d[A-Z]{2}\d\b"),
    re.compile(r"\b[A-Z]{2,}(?:-[A-Z0-9]+){1,}\b"),
    re.compile(r"\b\d{7,}\b"),
    re.compile(r"(?i)(?:(?<=\bde\s)|(?<=\baos\s)|(?<=\btem\s)|(?<=\baged\s))\d{1,3}\s+(?:anos|years)\b"),
    re.compile(r"(?i)\bBRCA[12]\b"),
    re.compile(r"(?i)\bHIV\s+positivo(?:\s+desde\s+\d{4})?"),
]


@dataclass
class Context:
    names: set[str] = field(default_factory=set)
    count: int = 0


def _mask(text: str, ctx: Context) -> str:
    def repl(m: re.Match) -> str:
        ctx.count += 1
        return "*"

    return re.sub(r"\S+", repl, text)


def _tokens(value: str) -> set[str]:
    value = re.sub(r"\([^)]*\)", " ", value)
    return {t for t in re.findall(CAP, value) if len(t) >= 3 and t not in STOP}


def learn_names(lines: list[str]) -> Context:
    ctx = Context()
    for line in lines:
        label, sep, value = line.partition(":")
        low = label.strip().lower()
        if sep and any(low.startswith(k) for k in NAME_LABELS + CHILD_LABELS):
            ctx.names |= _tokens(value)
        for rx in (TITLE_RE, PATIENT_RE):
            for m in rx.finditer(line):
                ctx.names |= _tokens(m.group(1))
    return ctx


def anonymize_line(line: str, ctx: Context) -> str:
    label, sep, value = line.partition(":")
    low = label.strip().lower()
    if sep and value.strip() and len(low) < 45 and any(k in low for k in LABEL_KEYWORDS):
        lead = value[: len(value) - len(value.lstrip())]
        return f"{label}:{lead}{_mask(value.strip(), ctx)}"

    for rx in PATTERNS:
        line = rx.sub(lambda m: _mask(m.group(0), ctx), line)
    for rx in (TITLE_RE, PATIENT_RE):
        line = rx.sub(lambda m: m.group(0).replace(m.group(1), _mask(m.group(1), ctx)), line)
    if ctx.names:
        names_rx = r"\b(?:" + "|".join(map(re.escape, sorted(ctx.names))) + r")\b"
        line = re.sub(names_rx, lambda m: _mask(m.group(0), ctx), line)
    return line


def _mask_terms(lines: list[str], terms: list[str]) -> tuple[list[str], int]:
    """Mask every occurrence of the Gemini-detected terms, one '*' per word."""
    clean = sorted({" ".join(t.split()) for t in terms if t.strip()}, key=len, reverse=True)
    if not clean:
        return list(lines), 0
    body = "|".join(r"\s+".join(map(re.escape, t.split())) for t in clean)
    rx = re.compile(rf"(?<!\w)(?:{body})(?!\w)", re.IGNORECASE)
    count = 0

    def repl(m: re.Match) -> str:
        nonlocal count
        words = len(m.group(0).split())
        count += words
        return " ".join("*" * words)

    return [rx.sub(repl, ln) for ln in lines], count


def censor_lines(lines: list[str]) -> tuple[list[str], int, str]:
    """Censor lines with Gemini; fall back to the built-in rules if it is unavailable."""
    terms = llm.detect_sensitive("\n".join(lines))
    if terms is not None:
        out, count = _mask_terms(lines, terms)
        return out, count, "gemini"
    ctx = learn_names(lines)
    return [anonymize_line(ln, ctx) for ln in lines], ctx.count, "rules"


def _read_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf_lines(doc) -> list[str]:
    lines = []
    for page in doc:
        lines += [ln.strip() for ln in page.get_text("text").splitlines() if ln.strip()]
    return lines


EXTRA_PATTERNS = (
    re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}\b"),
    re.compile(r"\b(?:[0-9A-Fa-f]{1,4}:){4,7}[0-9A-Fa-f]{1,4}\b"),
    re.compile(r"\b\d(?:[ \u00a0]\d){6,}\b"),
    re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"),
)


def _regex_terms(lines: list[str]) -> list[str]:
    """Deterministic safety net for IPs, e-mails and digit runs split by spaces."""
    text = "\n".join(lines)
    return [m.group(0) for rx in EXTRA_PATTERNS for m in rx.finditer(text)
            if not m.group(0).endswith("example") and "empresa" not in m.group(0).lower()]


def _detect_terms(lines: list[str]) -> tuple[list[str], str]:
    """Return the sensitive strings to hide (Gemini, or rules as a fallback)."""
    terms = llm.detect_sensitive("\n".join(lines))
    if terms is not None:
        return terms + _regex_terms(lines), "gemini"
    ctx = learn_names(lines)
    found = []
    for line in lines:
        masked = anonymize_line(line, ctx).split()
        for orig, new in zip(line.split(), masked):
            if new == "*" and orig != "*":
                found.append(orig)
    return found + _regex_terms(lines), "rules"


def _norm(word: str) -> str:
    plain = unicodedata.normalize("NFKD", word.strip(".,;:()<>[]\"'"))
    return "".join(c for c in plain if not unicodedata.combining(c)).lower()


def _redact_pdf(doc, terms: list[str]) -> None:
    """Redact in place: layout, tables and fonts stay; each hidden word becomes '*'."""
    tokens = sorted({tuple(_norm(w) for w in t.split()) for t in terms if t.strip()},
                    key=len, reverse=True)
    for page in doc:
        words = page.get_text("words")
        norm = [_norm(w[4]) for w in words]
        hidden = set()
        for tok in tokens:
            n = len(tok)
            if not all(tok):
                continue
            for i in range(len(words) - n + 1):
                if tuple(norm[i:i + n]) == tok:
                    hidden.update(range(i, i + n))
        stars = []
        spans = []
        for i in sorted(hidden):
            x0, y0, x1, y1 = words[i][:4]
            h = y1 - y0
            # Word boxes of tight lines overlap their neighbours, so only target the middle band.
            rect = pymupdf.Rect(x0 + 0.4, y0 + h * 0.3, x1 - 0.4, y1 - h * 0.3)
            page.add_redact_annot(rect, fill=False)
            stars.append((x0, x1, y1, h))
            prev = spans[-1] if spans else None
            if prev and abs(prev[1] - y0) < 2 and 0 <= x0 - prev[2] < 12:
                prev[2] = x1
            else:
                spans.append([x0, y0, x1, y1])
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)
        # A light bar over each run of hidden words closes the white gaps between the stars.
        for x0, y0, x1, y1 in spans:
            h = y1 - y0
            page.draw_rect(pymupdf.Rect(x0 - 0.5, y0 + h * 0.2, x1 + 0.5, y1 - h * 0.2),
                           color=None, fill=(0.82, 0.82, 0.82))
        for x0, x1, y1, h in stars:
            size = max(6, h * 0.85)
            width = pymupdf.get_text_length("*", fontname="helv", fontsize=size)
            page.insert_text(((x0 + x1 - width) / 2, y1 - h * 0.15), "*", fontname="helv", fontsize=size, color=(0, 0, 0))


def _process_pdf(data: bytes) -> tuple[bytes, str, str, int, str]:
    doc = pymupdf.open(stream=data, filetype="pdf")
    lines = _pdf_lines(doc)
    terms, engine = _detect_terms(lines)
    censored, count = _mask_terms(lines, terms)
    _redact_pdf(doc, terms)
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out, "\n".join(lines), "\n".join(censored), count, engine


def _docx_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _process_docx(data: bytes) -> tuple[bytes, str, str, int, str]:
    doc = Document(io.BytesIO(data))
    paragraphs = list(_docx_paragraphs(doc))
    original = [p.text for p in paragraphs]
    censored, count, engine = censor_lines(original)
    for p, text, new in zip(paragraphs, original, censored):
        if new != text and p.runs:
            p.runs[0].text = new
            for run in p.runs[1:]:
                run.text = ""
    buf = io.BytesIO()
    doc.save(buf)
    return (buf.getvalue(), "\n".join(t for t in original if t.strip()),
            "\n".join(t for t in censored if t.strip()), count, engine)


MIME = {
    "pdf": "application/pdf",
    "txt": "text/plain; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def process_file(data: bytes, ext: str) -> dict:
    """Return the censored file (same format) plus original/censored text for preview."""
    if ext == "txt":
        original = _read_text(data)
        cen, count, engine = censor_lines(original.splitlines())
        censored = "\n".join(cen)
        out = censored.encode("utf-8")
    elif ext == "pdf":
        out, original, censored, count, engine = _process_pdf(data)
    elif ext == "docx":
        out, original, censored, count, engine = _process_docx(data)
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX or TXT.")
    return {"file": out, "mime": MIME[ext], "original": original,
            "censored": censored, "count": count, "engine": engine}
